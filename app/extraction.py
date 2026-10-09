"""Document intelligence pipeline. All extracted facts remain reviewer-controlled candidates."""
import os, re, tempfile
from . import util

DOC_TYPE_KEYWORDS = {
 'commercial_invoice':['invoice','commercial invoice','seller','buyer','total due'],
 'packing_list':['packing list','net weight','gross weight','packages','carton'],
 'bill_of_lading':['bill of lading','b/l','vessel','port of loading','consignee'],
 'customs':['customs declaration','mrn','customs office','tariff','hs code'],
 'emissions':['embedded emissions','specific emissions','tco2e','direct emissions','indirect emissions'],
 'supplier_declaration':['supplier declaration','declaration of origin','we hereby declare','producer declares'],
 'methodology':['methodology','monitoring plan','calculation approach','system boundary'],
 'verification':['verification report','verifier','accredited','assurance statement']}
FACT_SPECS = [
 ('eu_tonnes',r'(?i)\\b(?:eu\\s+)?(?:shipment\\s+)?(?:quantity|net\\s+weight|weight|tonnes?|tons?)\\s*[:=]\\s*([\\d,]+(?:\\.\\d+)?)\\s*(?:t|tonnes?|tons?)?','t'),
 ('direct_intensity',r'(?i)\\b(?:direct|scope\\s*1)\\s*(?:emissions?\\s*)?(?:intensity)?\\s*[:=]\\s*([\\d,]+(?:\\.\\d+)?)','tCO2e/t'),
 ('indirect_intensity',r'(?i)\\b(?:indirect|scope\\s*2)\\s*(?:emissions?\\s*)?(?:intensity)?\\s*[:=]\\s*([\\d,]+(?:\\.\\d+)?)','tCO2e/t'),
 ('cn_code',r'\\b((?:25|28|31|72|73|76)\\d{6})\\b',''),
 ('invoice_number',r'(?i)\\binvoice\\s*(?:no\\.?|number|#)\\s*[:=-]?\\s*([A-Z0-9][A-Z0-9./_-]{2,})',''),
 ('supplier',r'(?im)^\\s*(?:supplier|manufacturer|vendor)\\s*[:=-]\\s*([^\\n|]{3,100})',''),
 ('shipment_reference',r'(?i)\\b(?:shipment|shipment\\s*ref|case\\s*ref)\\s*[:=-]\\s*([A-Z0-9._/-]{3,80})',''),
 ('origin_country',r'(?i)\\b(?:country\\s+of\\s+origin|origin)\\s*[:=-]\\s*([A-Za-z][A-Za-z .]{2,40})',''),
 ('installation',r'(?im)^\\s*(?:installation|plant|facility)\\s*[:=-]\\s*([^\\n|]{3,120})','')]

def classify(filename: str, text: str):
    low=(text or '').lower(); best,best_score='other',0
    for dtype,kws in DOC_TYPE_KEYWORDS.items():
        score=sum(1 for k in kws if k in low)
        if score>best_score: best,best_score=dtype,score
    return best,round(min(1.0,best_score/3) if best_score else 0.0,2)

def extract_file(filename: str, data: bytes):
    ext=os.path.splitext(filename.lower())[1]; text=''; meta={'type':ext,'warnings':[],'cells':[],'pages':[]}
    try:
        if ext in ('.txt','.csv'):
            text=data.decode('utf-8','replace')[:700000]
            for i,line in enumerate(text.splitlines()[:15000],1): meta['cells'].append({'location':f'line {i}','value':line[:10000]})
        elif ext=='.pdf':
            from pypdf import PdfReader
            with tempfile.NamedTemporaryFile(suffix='.pdf') as f:
                f.write(data); f.flush(); r=PdfReader(f.name)
                for i,p in enumerate(r.pages,1):
                    pt=p.extract_text() or ''
                    if len(pt.strip())<20:
                        try:
                            import fitz, pytesseract
                            from PIL import Image
                            page=fitz.open(f.name)[i-1]; pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
                            ocr=pytesseract.image_to_string(Image.frombytes('RGB',[pix.width,pix.height],pix.samples)) or ''
                            if ocr.strip(): pt=ocr; meta['warnings'].append(f'page {i}: OCR fallback used')
                        except Exception: meta['warnings'].append(f'page {i}: OCR unavailable in this environment')
                    meta['pages'].append({'page':i,'text':pt[:30000]}); text+=f'PAGE {i}\\n{pt}\\n'
                text=text[:900000]
        elif ext=='.xlsx':
            from openpyxl import load_workbook
            with tempfile.NamedTemporaryFile(suffix='.xlsx') as f:
                f.write(data); f.flush(); wb=load_workbook(f.name,data_only=True,read_only=True); chunks=[]
                for ws in wb.worksheets:
                    chunks.append('SHEET '+ws.title)
                    for row in ws.iter_rows():
                        vals=[]
                        for cell in row:
                            if cell.value is not None:
                                val=str(cell.value); vals.append(val); meta['cells'].append({'sheet':ws.title,'cell':cell.coordinate,'value':val[:10000]})
                        if vals: chunks.append(' | '.join(vals))
                text='\\n'.join(chunks)[:900000]
        else: meta['warnings'].append('Unsupported extractor. File fingerprint retained; no facts were promoted.')
    except Exception as e: meta['warnings'].append(f'extraction error: {e}')
    return text,meta

def infer_facts(text: str, meta: dict):
    found=[]
    for key,pat,unit in FACT_SPECS:
        m=re.search(pat,text)
        if not m: continue
        val=m.group(1).strip().strip('.,;'); loc,excerpt='text match',''
        for cell in meta.get('cells',[]):
            if val.lower() in str(cell.get('value','')).lower():
                loc=f"{cell.get('sheet','')+'!' if cell.get('sheet') else ''}{cell.get('cell',cell.get('location',''))}"; excerpt=str(cell.get('value',''))[:500]; break
        if loc=='text match':
            for pg in meta.get('pages',[]):
                if val.lower() in pg['text'].lower():
                    loc=f"page {pg['page']}"; excerpt=next((ln.strip() for ln in pg['text'].splitlines() if val.lower() in ln.lower()),'')[:500]; break
        if not excerpt:
            excerpt=' '.join(text[max(0,m.start()-120):min(len(text),m.end()+120)].split())[:500]
        found.append((key,val,unit,loc,excerpt,'high' if loc!='text match' else 'medium'))
    return found
