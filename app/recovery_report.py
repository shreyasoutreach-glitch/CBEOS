"""Buyer-facing evidence recovery workbook.

Exports a tenant-scoped, case-specific operational handoff. It intentionally
excludes raw extracted document text, credentials, and liability calculations.
The workbook reports workflow state; it does not certify CBAM compliance.
"""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe(value):
    """Keep untrusted strings as literal spreadsheet text, not formulas."""
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _owner_role(category, scope_type="case"):
    roles = {
        "commercial": "Procurement / accounts payable",
        "invoice": "Procurement / accounts payable",
        "shipping": "Logistics / customs broker",
        "classification": "Customs broker / trade compliance",
        "customs": "Customs broker / trade compliance",
        "installation": "Supplier plant contact",
        "activity": "Supplier plant / operations",
        "emissions": "Supplier sustainability / plant",
        "energy": "Supplier energy / plant",
        "methodology": "CBAM compliance owner",
        "carbon_price": "Finance / tax",
        "free_allocation": "CBAM compliance owner",
        "verification": "Supplier plant + accredited verifier",
        "supplier": "Procurement / supplier relationship owner",
    }
    return roles.get(str(category or "").casefold(),
                     "CBAM compliance owner" if scope_type == "case" else "Case owner")


def _next_action(category):
    actions = {
        "commercial": "Attach the invoice and reconcile quantity, product description and invoice reference.",
        "invoice": "Attach the invoice and reconcile quantity, product description and invoice reference.",
        "shipping": "Request the bill of lading, packing list or shipment record and match it to the import line.",
        "classification": "Have the customs broker confirm the applicable CN/TARIC code and record the source used.",
        "customs": "Attach the customs declaration and reconcile the declared code, origin and net mass.",
        "installation": "Confirm the producing installation identity and link supporting supplier evidence.",
        "activity": "Request the activity/production records and reconcile period, unit and production volume.",
        "emissions": "Request the installation emissions calculation and its underlying activity data.",
        "energy": "Request electricity/energy activity records and document the calculation boundary.",
        "methodology": "Confirm the applicable calculation method and document why it applies to this product and period.",
        "carbon_price": "Request evidence of carbon price legally paid and any rebate or compensation.",
        "free_allocation": "Review the applicable benchmark and free-allocation adjustment with a qualified reviewer.",
        "verification": "Request the installation-level verification report and confirm verifier status and scope.",
        "supplier": "Confirm the accountable supplier contact and send a scoped evidence request.",
    }
    return actions.get(str(category or "").casefold(),
                       "Confirm applicability, assign an owner and due date, then attach source evidence.")


def _rows(c, sql, args):
    return [dict(row) for row in c.execute(sql, args).fetchall()]


def build_recovery_workbook(c, tenant_id, case_id):
    case_rows = _rows(
        c, "SELECT * FROM cases WHERE tenant_id=? AND id=?", (tenant_id, case_id)
    )
    if not case_rows:
        return None
    case = case_rows[0]

    lines = _rows(c, """
        SELECT l.id, l.case_id, l.import_id, l.cn_code, l.description,
               l.origin_country, l.quantity, l.quantity_unit, l.invoice_ref,
               l.applicability, l.status, l.production_route,
               COALESCE(s.name, '') AS supplier_name,
               COALESCE(i.name, '') AS installation_name,
               COALESCE(im.reference, '') AS import_reference
        FROM import_lines l
        LEFT JOIN suppliers s ON s.id=l.supplier_id AND s.tenant_id=l.tenant_id
        LEFT JOIN installations i ON i.id=l.installation_id AND i.tenant_id=l.tenant_id
        LEFT JOIN imports im ON im.id=l.import_id AND im.tenant_id=l.tenant_id
        WHERE l.tenant_id=? AND l.case_id=? ORDER BY l.id
    """, (tenant_id, case_id))
    gaps = _rows(c, """
        SELECT id, name, category, applicability, status, scope_type, scope_id,
               reason, updated
        FROM evidence_requirements
        WHERE tenant_id=? AND case_id=? AND status='missing'
          AND applicability!='not_applicable'
        ORDER BY CASE scope_type WHEN 'case' THEN 0 WHEN 'import_line' THEN 1
                 WHEN 'supplier' THEN 2 ELSE 3 END, scope_id, name
    """, (tenant_id, case_id))
    exceptions = _rows(c, """
        SELECT id, title, category, severity, status, affected_entity_type,
               affected_entity_id, detail, field, recommended_action, owner,
               deadline, source, updated
        FROM exceptions
        WHERE tenant_id=? AND case_id=?
          AND status NOT IN ('RESOLVED','ACCEPTED_WITH_RISK','WAIVED')
        ORDER BY CASE severity WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END, id
    """, (tenant_id, case_id))
    requests = _rows(c, """
        SELECT r.id, s.name AS supplier_name, r.title, r.requirement_name,
               r.status, r.deadline, r.sent_at, r.responded_at,
               r.escalation_level, r.draft
        FROM supplier_requests r
        JOIN suppliers s ON s.id=r.supplier_id AND s.tenant_id=r.tenant_id
        WHERE r.tenant_id=? AND r.case_id=? ORDER BY
          CASE r.status WHEN 'OVERDUE' THEN 0 WHEN 'DRAFT' THEN 1 ELSE 2 END, r.id
    """, (tenant_id, case_id))
    documents = _rows(c, """
        SELECT d.id, d.filename, d.doc_type, d.status, d.scan_status,
               d.size_bytes, d.sha256, d.version, d.uploaded,
               d.import_line_id, d.supplier_id, d.installation_id,
               COALESCE(s.name, '') AS supplier_name,
               COALESCE(i.name, '') AS installation_name,
               (SELECT COUNT(*) FROM facts f
                WHERE f.tenant_id=d.tenant_id AND f.document_id=d.id) AS candidate_fact_count
        FROM documents d
        LEFT JOIN suppliers s ON s.id=d.supplier_id AND s.tenant_id=d.tenant_id
        LEFT JOIN installations i ON i.id=d.installation_id AND i.tenant_id=d.tenant_id
        WHERE d.tenant_id=? AND d.case_id=? ORDER BY d.id
    """, (tenant_id, case_id))

    open_count = len(exceptions)
    overdue_count = sum(1 for r in requests if str(r["status"]).upper() == "OVERDUE")
    wb = Workbook()
    summary = wb.active
    summary.title = "Read Me & Summary"
    summary_rows = [
        ("CBEOS | Evidence Recovery Handoff", ""),
        ("Case", case.get("case_name", "")),
        ("Company", case.get("company", "")),
        ("Reporting period", case.get("period", "")),
        ("Sector", case.get("sector", "")),
        ("Report generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Import lines", len(lines)),
        ("Missing evidence requirements", len(gaps)),
        ("Open exceptions", open_count),
        ("Supplier requests overdue", overdue_count),
        ("Documents in register", len(documents)),
        ("Purpose", "Actionable operational handoff: missing evidence, unresolved exceptions, supplier follow-up and document traceability."),
        ("Use boundary", "Workflow status only. Not a legal opinion, compliance certification, verified emissions calculation, declaration, or liability figure."),
        ("Data handling", "Contains case metadata and document register metadata. Share only with authorized case participants."),
        ("Recommended next step", "Assign an owner and due date to every open gap and exception; confirm source evidence with a qualified reviewer."),
    ]
    for row in summary_rows:
        summary.append([_safe(v) for v in row])

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append([_safe(h) for h in headers])
        for row in rows:
            ws.append([_safe(v) for v in row])
        return ws

    sheet("Import Lines", [
        "Line ID", "Import reference", "CN code", "Description", "Origin country",
        "Quantity", "Unit", "Supplier", "Installation", "Production route",
        "Invoice reference", "Applicability", "Workflow status"
    ], [[r["id"], r["import_reference"], r["cn_code"], r["description"],
          r["origin_country"], r["quantity"], r["quantity_unit"], r["supplier_name"],
          r["installation_name"], r["production_route"], r["invoice_ref"],
          r["applicability"], r["status"]] for r in lines])

    sheet("Evidence Gaps", [
        "Requirement ID", "Requirement", "Category", "Applicability", "Status",
        "Scope type", "Scope ID", "Suggested owner role", "Suggested next action",
        "Why / review note", "Last updated"
    ], [[r["id"], r["name"], r["category"], r["applicability"], r["status"],
          r["scope_type"], r["scope_id"], _owner_role(r["category"], r["scope_type"]),
          _next_action(r["category"]), r["reason"], r["updated"]] for r in gaps])

    sheet("Open Exceptions", [
        "Exception ID", "Title", "Category", "Severity", "Status", "Entity type",
        "Entity ID", "Detail", "Field", "Suggested owner role", "Recommended next action", "Owner",
        "Deadline", "Source", "Last updated"
    ], [[r["id"], r["title"], r["category"], r["severity"], r["status"],
          r["affected_entity_type"], r["affected_entity_id"], r["detail"],
          r["field"], _owner_role(r["category"], r["affected_entity_type"] or "case"),
          r["recommended_action"] or _next_action(r["category"]), r["owner"], r["deadline"],
          r["source"], r["updated"]] for r in exceptions])

    sheet("Supplier Follow-up", [
        "Request ID", "Supplier", "Request title", "Requirement", "Status",
        "Deadline", "Sent at", "Responded at", "Escalation level", "Draft text"
    ], [[r["id"], r["supplier_name"], r["title"], r["requirement_name"],
          r["status"], r["deadline"], r["sent_at"], r["responded_at"],
          r["escalation_level"], r["draft"]] for r in requests])

    sheet("Document Register", [
        "Document ID", "Filename", "Document type", "Status", "Scan status",
        "Size bytes", "SHA-256", "Version", "Uploaded at", "Import line ID",
        "Supplier", "Installation", "Candidate facts (count)"
    ], [[r["id"], r["filename"], r["doc_type"], r["status"], r["scan_status"],
          r["size_bytes"], r["sha256"], r["version"], r["uploaded"],
          r["import_line_id"], r["supplier_name"], r["installation_name"],
          r["candidate_fact_count"]] for r in documents])

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2" if ws.title != "Read Me & Summary" else "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if ws.title == "Read Me & Summary":
            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 105
            for cell in ws["A"]:
                cell.font = Font(bold=True, color="17365D")
        else:
            for col_idx in range(1, ws.max_column + 1):
                values = [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, min(ws.max_row, 100) + 1)]
                width = min(max(max(values or [12]) + 2, 12), 42)
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            ws.auto_filter.ref = ws.dimensions

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
