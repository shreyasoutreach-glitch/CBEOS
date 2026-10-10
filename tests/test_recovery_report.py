"""Buyer deliverable tests for the case-scoped recovery workbook."""
import sqlite3
import sys
import unittest
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.recovery_report import build_recovery_workbook


def fixture_db():
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    c.executescript("""
    CREATE TABLE cases(id INTEGER, tenant_id INTEGER, company TEXT, case_name TEXT, period TEXT, sector TEXT, status TEXT);
    CREATE TABLE import_lines(id INTEGER, tenant_id INTEGER, case_id INTEGER, import_id INTEGER, cn_code TEXT, description TEXT, origin_country TEXT, quantity TEXT, quantity_unit TEXT, invoice_ref TEXT, applicability TEXT, status TEXT, production_route TEXT, supplier_id INTEGER, installation_id INTEGER);
    CREATE TABLE suppliers(id INTEGER, tenant_id INTEGER, name TEXT);
    CREATE TABLE installations(id INTEGER, tenant_id INTEGER, name TEXT);
    CREATE TABLE imports(id INTEGER, tenant_id INTEGER, reference TEXT);
    CREATE TABLE evidence_requirements(id INTEGER, tenant_id INTEGER, case_id INTEGER, name TEXT, category TEXT, applicability TEXT, status TEXT, scope_type TEXT, scope_id INTEGER, reason TEXT, updated TEXT);
    CREATE TABLE exceptions(id INTEGER, tenant_id INTEGER, case_id INTEGER, title TEXT, category TEXT, severity TEXT, status TEXT, affected_entity_type TEXT, affected_entity_id INTEGER, detail TEXT, field TEXT, recommended_action TEXT, owner TEXT, deadline TEXT, source TEXT, updated TEXT);
    CREATE TABLE supplier_requests(id INTEGER, tenant_id INTEGER, case_id INTEGER, supplier_id INTEGER, title TEXT, requirement_name TEXT, status TEXT, deadline TEXT, sent_at TEXT, responded_at TEXT, escalation_level INTEGER, draft TEXT);
    CREATE TABLE documents(id INTEGER, tenant_id INTEGER, case_id INTEGER, filename TEXT, doc_type TEXT, status TEXT, scan_status TEXT, size_bytes INTEGER, sha256 TEXT, version INTEGER, uploaded TEXT, import_line_id INTEGER, supplier_id INTEGER, installation_id INTEGER);
    CREATE TABLE facts(id INTEGER, tenant_id INTEGER, document_id INTEGER);
    """)
    c.execute("INSERT INTO cases VALUES(1, 7, 'Acme Imports', '=CMD(1)', '2026', 'iron_steel', 'working')")
    c.execute("INSERT INTO suppliers VALUES(10, 7, 'Steel Supplier')")
    c.execute("INSERT INTO installations VALUES(20, 7, 'Mill A')")
    c.execute("INSERT INTO imports VALUES(30, 7, 'SHIP-001')")
    c.execute("INSERT INTO import_lines VALUES(40, 7, 1, 30, '72081000', 'Coil', 'India', '12.5', 't', 'INV-1', 'needs_review', 'blocked', 'BF-BOF', 10, 20)")
    c.execute("INSERT INTO evidence_requirements VALUES(50, 7, 1, 'Verified emissions', 'verification', 'applicable', 'missing', 'case', 0, 'Need verifier evidence', '2026-10-11')")
    c.execute("INSERT INTO exceptions VALUES(60, 7, 1, 'Missing verifier report', 'verification', 'high', 'OPEN', 'import_line', 40, 'No report attached', 'verification', 'Request report from supplier', 'Owner A', '2026-10-20', 'document review', '2026-10-11')")
    c.execute("INSERT INTO supplier_requests VALUES(70, 7, 1, 10, 'Send report', 'verification', 'OVERDUE', '2026-10-10', '', '', 1, 'Please send report')")
    c.execute("INSERT INTO documents VALUES(80, 7, 1, 'report.pdf', 'verification', 'processed', 'clean', 1234, 'abc123', 1, '2026-10-11', 40, 10, 20)")
    c.execute("INSERT INTO facts VALUES(90, 7, 80)")
    return c


class RecoveryWorkbookTests(unittest.TestCase):
    def test_workbook_contains_actionable_sheets_and_scoped_records(self):
        c = fixture_db()
        payload = build_recovery_workbook(c, 7, 1)
        self.assertIsInstance(payload, bytes)
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, [
            "Read Me & Summary", "Import Lines", "Evidence Gaps",
            "Open Exceptions", "Supplier Follow-up", "Pilot Measurement", "Document Register"
        ])
        summary = list(wb["Read Me & Summary"].values)
        self.assertIn(("Missing evidence requirements", 1), summary)
        self.assertIn(("Open exceptions", 1), summary)
        self.assertIn(("Supplier requests overdue", 1), summary)
        self.assertEqual(wb["Import Lines"].max_row, 2)
        self.assertEqual(wb["Evidence Gaps"].max_row, 2)
        self.assertEqual(wb["Open Exceptions"].max_row, 2)
        gap_headers = [cell.value for cell in wb["Evidence Gaps"][1]]
        self.assertIn("Suggested owner role", gap_headers)
        self.assertIn("Suggested next action", gap_headers)
        self.assertIn("Suggested owner role", [cell.value for cell in wb["Open Exceptions"][1]])
        self.assertIn("Request the installation-level verification report",
                      wb["Evidence Gaps"].cell(row=2, column=9).value)
        self.assertIn("Request the installation-level verification report",
                      wb["Open Exceptions"].cell(row=2, column=11).value)
        self.assertEqual(wb["Supplier Follow-up"].max_row, 2)
        self.assertEqual(wb["Document Register"].max_row, 2)
        # Untrusted text must remain a literal string, not an Excel formula.
        self.assertEqual(wb["Read Me & Summary"]["B2"].value, "'=CMD(1)")
        self.assertEqual(wb["Read Me & Summary"]["B2"].data_type, "s")
        self.assertEqual(wb["Read Me & Summary"]["B3"].value, "Acme Imports")
        self.assertEqual(wb["Read Me & Summary"]["B3"].data_type, "s")
        self.assertNotIn("No legal", str(list(wb["Read Me & Summary"].values)))
        wb.close()
        c.close()

    def test_report_does_not_cross_tenant_boundary(self):
        c = fixture_db()
        self.assertIsNone(build_recovery_workbook(c, 8, 1))
        c.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
